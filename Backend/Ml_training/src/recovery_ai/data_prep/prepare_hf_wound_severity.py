from __future__ import annotations

import argparse
import re
from pathlib import Path
from io import BytesIO

from openpyxl import load_workbook
import pandas as pd
from PIL import Image


RISK_MAP = {
    "mild": "normal",
    "mild/moderate": "mild_concern",
    "moderate": "mild_concern",
    "moderate to severe": "urgent",
    "severe": "urgent",
}


def normalize_severity(value: object, messages: object) -> str | None:
    text = str(value or "").strip().lower()
    if text in RISK_MAP:
        return text

    message_text = str(messages).lower()
    patterns = [
        (r'"injury-severity"\s*:\s*"([^"]+)"', message_text),
        (r"injury-severity['\"]?\s*:\s*['\"]?([a-z /-]+)", message_text),
        (r"\b(moderate to severe|mild/moderate|severe|moderate|mild)\b", message_text),
    ]
    for pattern, source in patterns:
        match = re.search(pattern, source)
        if not match:
            continue
        candidate = match.group(1).strip().lower().rstrip(".")
        if candidate in RISK_MAP:
            return candidate
        if "moderate to severe" in candidate:
            return "moderate to severe"
        if "severe" in candidate:
            return "severe"
        if "moderate" in candidate:
            return "moderate"
        if "mild" in candidate:
            return "mild"
    return None


def image_from_cell(cell: object) -> Image.Image:
    if cell is None:
        raise ValueError("Missing image")
    if isinstance(cell, Image.Image):
        return cell.convert("RGB")
    if isinstance(cell, dict) and "bytes" in cell:
        from io import BytesIO

        return Image.open(BytesIO(cell["bytes"])).convert("RGB")
    raise TypeError(f"Unsupported image cell type: {type(cell)!r}")


def prepare_dataset(raw_dir: Path, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    for class_name in sorted(set(RISK_MAP.values())):
        (output_dir / class_name).mkdir(parents=True, exist_ok=True)

    rows: list[dict[str, str]] = []
    parquet_files = sorted((raw_dir / "data").glob("*.parquet"))
    workbook_path = raw_dir / "finedata.xlsx"
    if not parquet_files and not workbook_path.exists():
        raise FileNotFoundError(f"No parquet files found under {raw_dir / 'data'}")

    counter = {class_name: 0 for class_name in sorted(set(RISK_MAP.values()))}
    skipped = 0
    for parquet_path in parquet_files:
        df = pd.read_parquet(parquet_path)
        for index, row in df.iterrows():
            severity = normalize_severity(row.get("injury-severity"), row.get("messages"))
            if severity is None:
                skipped += 1
                continue
            risk_class = RISK_MAP[severity]
            try:
                image = image_from_cell(row["image"])
            except (TypeError, ValueError):
                skipped += 1
                continue
            counter[risk_class] += 1
            filename = f"{risk_class}_{counter[risk_class]:05d}.jpg"
            output_path = output_dir / risk_class / filename
            image.save(output_path, quality=92)
            rows.append(
                {
                    "source": str(parquet_path),
                    "source_index": str(index),
                    "filename": str(output_path),
                    "severity": severity,
                    "risk_class": risk_class,
                    "note": "Weak label parsed from public accident/wound instruction dataset.",
                }
            )

    if workbook_path.exists():
        workbook = load_workbook(workbook_path, read_only=False)
        sheet = workbook[workbook.sheetnames[0]]
        images_by_row = {}
        for image_obj in getattr(sheet, "_images", []):
            row_number = image_obj.anchor._from.row + 1
            images_by_row[row_number] = image_obj

        for row_number in range(2, sheet.max_row + 1):
            text = sheet.cell(row_number, 3).value
            image_obj = images_by_row.get(row_number)
            severity = normalize_severity("Unknown", text)
            if severity is None or image_obj is None:
                skipped += 1
                continue
            risk_class = RISK_MAP[severity]
            image = Image.open(BytesIO(image_obj._data())).convert("RGB")
            counter[risk_class] += 1
            filename = f"{risk_class}_{counter[risk_class]:05d}.jpg"
            output_path = output_dir / risk_class / filename
            image.save(output_path, quality=92)
            rows.append(
                {
                    "source": str(workbook_path),
                    "source_index": str(row_number),
                    "filename": str(output_path),
                    "severity": severity,
                    "risk_class": risk_class,
                    "note": "Weak label parsed from public accident/wound Excel dataset.",
                }
            )

    manifest_path = output_dir / "manifest.csv"
    pd.DataFrame(rows).to_csv(manifest_path, index=False)
    print(f"Wrote {len(rows)} images to {output_dir}")
    print(f"Skipped {skipped} rows without usable severity labels")
    print("Class counts:", counter)
    print(f"Manifest: {manifest_path}")
    return manifest_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-dir", default="data/raw/hf_accidents_wounds_finetune")
    parser.add_argument("--output-dir", default="data/processed/classification")
    args = parser.parse_args()
    prepare_dataset(Path(args.raw_dir), Path(args.output_dir))


if __name__ == "__main__":
    main()
